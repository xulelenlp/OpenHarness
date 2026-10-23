"""Word 报告工具测试。"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING

from openharness.tools.base import ToolExecutionContext
from qaapp.tools.report import (
    ReportSection,
    ReportTable,
    WriteReportDocxInput,
    WriteReportDocxTool,
    parse_markdown,
)
from qaapp.tools.script import CONVERSATION_DIR_KEY


def _context(tmp_path: Path) -> ToolExecutionContext:
    return ToolExecutionContext(cwd=tmp_path, metadata={CONVERSATION_DIR_KEY: str(tmp_path)})


@pytest.mark.asyncio
async def test_generate_docx_from_markdown(tmp_path: Path):
    markdown = """# 一、总体情况

2026年9月福建售电量为 **64.9459 亿千瓦时**，同比下降 6.31%。

## 1.1 分产业

| 产业 | 售电量 | 同比 |
|---|---:|---:|
| 第一产业 | 1.1044 | -13.85 |
| 第二产业 | 36.7980 | 0.20 |

- 第一产业降幅最大
- 第二产业基本持平
"""
    tool = WriteReportDocxTool()
    result = await tool.execute(
        WriteReportDocxInput(
            title="福建售电量分析",
            org="福建",
            period="2026-09",
            markdown=markdown,
            data_quality=["电费回收率序列恒定"],
            sources=[{"label": "业务库", "detail": "zhss_result_data_qs"}],
        ),
        _context(tmp_path),
    )
    assert result.is_error is False
    path = Path(result.metadata["path"])
    assert path.is_file()
    doc = Document(str(path))
    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert any("总体情况" in text for text in texts)
    assert len(doc.tables) == 1
    assert any("数据质量与局限性" in text for text in texts)
    # 默认标题/单位/日期与正文同页，不额外生成目录页。
    assert all(text != "目  录" for text in texts)


@pytest.mark.asyncio
async def test_generate_docx_from_sections(tmp_path: Path):
    tool = WriteReportDocxTool()
    result = await tool.execute(
        WriteReportDocxInput(
            title="结构化报告",
            sections=[
                ReportSection(
                    heading="一、总体情况",
                    paragraphs=["这是正文。"],
                    bullets=["要点一"],
                    tables=[ReportTable(columns=["指标", "值"], rows=[["售电量", "64.9459"]])],
                )
            ],
        ),
        _context(tmp_path),
    )
    assert result.is_error is False
    doc = Document(result.metadata["path"])
    assert len(doc.tables) == 1
    assert any("这是正文" in paragraph.text for paragraph in doc.paragraphs)


@pytest.mark.asyncio
async def test_path_traversal_rejected(tmp_path: Path):
    tool = WriteReportDocxTool()
    result = await tool.execute(
        WriteReportDocxInput(title="x", markdown_path="../../etc/passwd"),
        _context(tmp_path),
    )
    assert result.is_error is True
    assert "越界" in result.output


@pytest.mark.asyncio
async def test_missing_image_warns_but_still_generates(tmp_path: Path):
    tool = WriteReportDocxTool()
    result = await tool.execute(
        WriteReportDocxInput(
            title="缺图报告",
            sections=[ReportSection(heading="一、图", images=[{"path": "outputs/missing.png", "caption": "图1"}])],
        ),
        _context(tmp_path),
    )
    assert result.is_error is False
    assert "图片不存在" in "；".join(result.metadata["warnings"])


def test_parse_markdown_table_and_heading():
    blocks = parse_markdown("# 标题\n\n| A | B |\n|---|---|\n| 1 | 2 |\n")
    kinds = [block["kind"] for block in blocks]
    assert "heading" in kinds
    assert "table" in kinds


def _write_png(path: Path, width: int, height: int) -> None:
    """写一张纯白 PNG（避免为测试引入 Pillow 依赖）。"""
    import struct
    import zlib

    raw = b"".join(b"\x00" + b"\xff\xff\xff" * width for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        crc = zlib.crc32(tag + data) & 0xFFFFFFFF
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", crc)

    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


@pytest.mark.asyncio
async def test_image_is_floating_top_and_bottom(tmp_path: Path):
    """图片是“上下环绕”浮动对象，且图片段落不额外占一整行高度。"""
    from docx.oxml.ns import qn

    image_path = tmp_path / "outputs" / "figure.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    _write_png(image_path, 1200, 600)

    result = await WriteReportDocxTool().execute(
        WriteReportDocxInput(
            title="图片版式",
            sections=[
                ReportSection(
                    heading="一、图",
                    paragraphs=["图前正文。"],
                    images=[{"path": "outputs/figure.png", "caption": "图 1 测试"}],
                )
            ],
        ),
        _context(tmp_path),
    )
    assert result.is_error is False
    doc = Document(result.metadata["path"])

    anchors = doc.element.body.findall(".//" + qn("wp:anchor"))
    assert len(anchors) == 1
    assert doc.element.body.findall(".//" + qn("wp:inline")) == []
    assert anchors[0].find(qn("wp:wrapTopAndBottom")) is not None
    # 竖直方向按段落顶部对齐（align=top），避免部分阅读器把图排到前一段之前。
    position_v = anchors[0].find(qn("wp:positionV"))
    assert position_v.get("relativeFrom") == "paragraph"
    assert position_v.find(qn("wp:align")).text == "top"

    picture_paragraph = next(p for p in doc.paragraphs if p._p.findall(".//" + qn("wp:anchor")))
    # 不能用 29pt 固定行距（会裁切图片），也不能按图高预留行高（会多出大段空行）。
    assert picture_paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
    assert picture_paragraph.paragraph_format.space_before.pt == 6
    assert picture_paragraph.paragraph_format.keep_with_next is True

    p_pr = picture_paragraph._p.find(qn("w:pPr"))
    assert p_pr.find(qn("w:snapToGrid")).get(qn("w:val")) == "0"
    assert p_pr.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")) == "2"

    caption = next(p for p in doc.paragraphs if p.text == "图 1 测试")
    assert caption.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
    assert caption.paragraph_format.space_before.pt == 6
    texts = [p.text for p in doc.paragraphs]
    assert "图前正文。" in texts


@pytest.mark.asyncio
async def test_multiple_images_stay_compact(tmp_path: Path):
    """连续两张图都要用单倍行距 + 1pt 段落标记，不能只处理第一张。"""
    from docx.oxml.ns import qn

    image_path = tmp_path / "outputs" / "figure.png"
    image_path.parent.mkdir(parents=True, exist_ok=True)
    _write_png(image_path, 1200, 600)

    result = await WriteReportDocxTool().execute(
        WriteReportDocxInput(
            title="多图版式",
            sections=[
                ReportSection(
                    heading="一、两张图",
                    images=[
                        {"path": "outputs/figure.png", "caption": "图 1"},
                        {"path": "outputs/figure.png", "caption": "图 2"},
                    ],
                )
            ],
        ),
        _context(tmp_path),
    )
    assert result.is_error is False
    doc = Document(result.metadata["path"])

    picture_paragraphs = [
        p for p in doc.paragraphs if p._p.findall(".//" + qn("wp:anchor"))
    ]
    assert len(picture_paragraphs) == 2
    for paragraph in picture_paragraphs:
        assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE
        p_pr = paragraph._p.find(qn("w:pPr"))
        assert p_pr.find(qn("w:rPr")).find(qn("w:sz")).get(qn("w:val")) == "2"
    captions = [p.text for p in doc.paragraphs if p.text in {"图 1", "图 2"}]
    assert captions == ["图 1", "图 2"]


@pytest.mark.asyncio
async def test_report_uses_state_grid_official_format(tmp_path: Path):
    """生成的 Word 必须符合国网公文格式（字体/字号/行距/页边距）。"""
    markdown = (
        "# 一、总体情况\n\n"
        "这是正文。\n\n"
        "## （二）分项标题\n\n"
        "### 3.三级标题\n\n"
        "#### （4）四级标题\n\n"
        "正文第二段。\n"
    )
    result = await WriteReportDocxTool().execute(
        WriteReportDocxInput(
            title="关于测试问题的汇报",
            org="国家电网公司",
            date="2026年10月21日",
            markdown=markdown,
            include_toc=False,
            include_page_numbers=False,
        ),
        _context(tmp_path),
    )
    assert result.is_error is False
    doc = Document(result.metadata["path"])

    section = doc.sections[0]
    assert round(section.top_margin.mm) == 35
    assert round(section.bottom_margin.mm) == 35
    assert round(section.left_margin.mm) == 28
    assert round(section.right_margin.mm) == 28

    normal = doc.styles["Normal"]
    assert normal.font.name == "仿宋"
    assert normal.font.size.pt == 16
    assert normal.paragraph_format.line_spacing.pt == 29
    assert normal.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY

    expected = {
        1: ("黑体", False),
        2: ("楷体", True),
        3: ("宋体", True),
        4: ("仿宋", False),
    }
    for level, (font, bold) in expected.items():
        style = doc.styles[f"Heading {level}"]
        assert style.font.name == font
        assert style.font.size.pt == 16
        assert bool(style.font.bold) is bold
        assert style.paragraph_format.line_spacing.pt == 29

    def _paragraph(text: str):
        return next(p for p in doc.paragraphs if p.text == text)

    title_run = _paragraph("关于测试问题的汇报").runs[0]
    assert title_run.font.name == "宋体"
    assert title_run.font.size.pt == 22
    assert title_run.font.bold is True

    for value in ("国家电网公司", "2026年10月21日"):
        run = _paragraph(value).runs[0]
        assert run.font.name == "楷体"
        assert run.font.size.pt == 16
        assert run.font.bold is True

    for text, level, font, bold in (
        ("一、总体情况", 1, "黑体", False),
        ("（二）分项标题", 2, "楷体", True),
        ("3.三级标题", 3, "宋体", True),
        ("（4）四级标题", 4, "仿宋", False),
    ):
        paragraph = _paragraph(text)
        assert paragraph.style.name == f"Heading {level}"
        run = paragraph.runs[0]
        assert run.font.name == font
        assert run.font.size.pt == 16
        assert bool(run.font.bold) is bold

    body = _paragraph("这是正文。")
    assert body.runs[0].font.name == "仿宋"
    assert body.runs[0].font.size.pt == 16
    assert body.paragraph_format.line_spacing.pt == 29
    assert body.paragraph_format.first_line_indent.pt == 32
