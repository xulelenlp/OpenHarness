"""文档解析：把用户上传的 Word / Excel / 文本文件转成可注入模型的纯文本。

设计要点：

- **主动解析**：上传即解析（不依赖模型自己读二进制文件），解析结果落盘为
  ``extracted/<原名>.txt``，模型需要细节时可用 ``read_file`` 读取该文件。
- **纯文本优先**：所有格式最终都归一成 UTF-8 文本，Word 表格与 Excel 工作表
  转成 Markdown/TSV，便于模型与后续脚本消费。
- **失败不丢文件**：解析异常（加密、损坏、旧版 .doc）只标记失败并给出原因，
  原始文件仍然保留，模型可以自行尝试其它处理方式。
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

#: 文本类文件一次最多读取多少字节（避免超大 txt 直接把内存吃满）。
MAX_TEXT_READ_BYTES = 8 * 1024 * 1024
#: 压缩包（docx/xlsx）解压后的总量上限与压缩比上限，防 zip 炸弹。
MAX_ZIP_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_ZIP_RATIO = 200

#: 后缀 → 种类（同时充当支持白名单）。
KIND_BY_SUFFIX: dict[str, str] = {
    ".txt": "text",
    ".log": "text",
    ".md": "markdown",
    ".csv": "csv",
    ".json": "json",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".xml": "xml",
    ".docx": "word",
    ".xlsx": "excel",
    ".xlsm": "excel",
    ".xls": "excel",
}

SUPPORTED_SUFFIXES: tuple[str, ...] = tuple(sorted(KIND_BY_SUFFIX))

TRUNCATION_NOTE = "\n\n...[解析内容已截断，完整内容见解析文本文件]"

#: 旧版二进制 Office 格式：明确不支持，给出「另存为」提示。
LEGACY_OFFICE_HINTS = {
    ".doc": "旧版 .doc 暂不支持，请在 Word 里另存为 .docx 后重新上传",
    ".ppt": "暂不支持 PPT，请导出为 .txt / .md 后上传",
    ".xlsb": "暂不支持 .xlsb，请另存为 .xlsx 后重新上传",
}


class UnsupportedDocument(Exception):
    """文件类型不支持或内容无法解析。"""


class DocumentTooLarge(Exception):
    """解析前的保护性拒绝（压缩炸弹等）。"""


@dataclass
class ParsedDocument:
    """一次解析的结果。"""

    kind: str
    text: str
    truncated: bool = False
    meta: dict[str, Any] = field(default_factory=dict)

    @property
    def chars(self) -> int:
        return len(self.text)


def suffix_of(filename: str) -> str:
    """返回小写后缀（含点）；无后缀返回空串。"""
    return Path(filename).suffix.lower()


def kind_for_suffix(suffix: str) -> str | None:
    return KIND_BY_SUFFIX.get(suffix.lower())


def parse_document(
    path: str | Path,
    *,
    max_chars: int,
    max_excel_rows_per_sheet: int = 2_000,
) -> ParsedDocument:
    """解析一个已落盘的上传文件。

    抛出 :class:`UnsupportedDocument` 表示类型不支持或内容损坏；调用方应把文件
    保留并把 ``parse_status`` 标记为 ``failed``，而不是丢弃用户文件。
    """
    file_path = Path(path)
    suffix = suffix_of(file_path.name)
    kind = kind_for_suffix(suffix)
    if kind is None:
        hint = LEGACY_OFFICE_HINTS.get(suffix)
        raise UnsupportedDocument(hint or f"暂不支持的文件类型：{suffix or '无后缀'}")

    if kind in {"text", "markdown", "json", "yaml", "xml"}:
        parsed = _parse_text(file_path, kind)
    elif kind == "csv":
        parsed = _parse_csv(file_path, max_excel_rows_per_sheet)
    elif kind == "word":
        parsed = _parse_docx(file_path)
    elif kind == "excel":
        parsed = _parse_excel(file_path, suffix, max_excel_rows_per_sheet)
    else:  # pragma: no cover - 白名单已穷举
        raise UnsupportedDocument(f"暂不支持的文件类型：{suffix}")

    if max_chars > 0 and len(parsed.text) > max_chars:
        parsed.text = parsed.text[:max_chars] + TRUNCATION_NOTE
        parsed.truncated = True
    return parsed


def save_extracted_text(path: str | Path, parsed: ParsedDocument) -> Path:
    """把解析文本写入 ``extracted/*.txt``（UTF-8）。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(parsed.text, encoding="utf-8")
    return target


# ---------------------------------------------------------------------------
# 文本类
# ---------------------------------------------------------------------------


def decode_bytes(data: bytes) -> tuple[str, str]:
    """按 utf-8 → gb18030 → latin-1 逐级尝试解码，返回 (文本, 编码名)。"""
    for encoding in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8?replace"


def _read_text_bytes(path: Path) -> bytes:
    with path.open("rb") as handle:
        return handle.read(MAX_TEXT_READ_BYTES)


def _parse_text(path: Path, kind: str) -> ParsedDocument:
    text, encoding = decode_bytes(_read_text_bytes(path))
    meta: dict[str, Any] = {"encoding": encoding, "lines": text.count("\n") + 1}
    if kind == "json":
        try:
            import json as _json

            parsed_json = _json.loads(text)
            meta["json_root"] = type(parsed_json).__name__
        except (ValueError, TypeError):
            meta["json_valid"] = False
    return ParsedDocument(kind=kind, text=text, meta=meta)


def _parse_csv(path: Path, max_rows: int) -> ParsedDocument:
    text, encoding = decode_bytes(_read_text_bytes(path))
    reader = csv.reader(io.StringIO(text))
    rows: list[list[str]] = []
    truncated = False
    for index, row in enumerate(reader):
        if max_rows > 0 and index >= max_rows:
            truncated = True
            break
        rows.append([cell.strip() for cell in row])
    rendered, columns = _rows_to_markdown(rows)
    if truncated:
        rendered += f"\n\n...[仅解析前 {max_rows} 行]"
    return ParsedDocument(
        kind="csv",
        text=rendered,
        truncated=truncated,
        meta={"encoding": encoding, "rows": len(rows), "columns": columns},
    )


def _rows_to_markdown(rows: list[list[str]]) -> tuple[str, int]:
    """二维表 → Markdown 表格；返回 (文本, 列数)。"""
    if not rows:
        return "(空表)", 0
    columns = max(len(row) for row in rows)
    normalized = [row + [""] * (columns - len(row)) for row in rows]
    header = normalized[0]
    body = normalized[1:]
    lines = ["| " + " | ".join(_escape_cell(cell) for cell in header) + " |"]
    lines.append("| " + " | ".join("---" for _ in range(columns)) + " |")
    for row in body:
        lines.append("| " + " | ".join(_escape_cell(cell) for cell in row) + " |")
    return "\n".join(lines), columns


def _escape_cell(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ").strip()


# ---------------------------------------------------------------------------
# Word（.docx）
# ---------------------------------------------------------------------------


def _guard_zip(path: Path) -> None:
    """检查压缩包规模，防止 zip 炸弹。"""
    try:
        with zipfile.ZipFile(path) as archive:
            total = sum(item.file_size for item in archive.infolist())
            compressed = sum(max(item.compress_size, 1) for item in archive.infolist())
    except (zipfile.BadZipFile, OSError) as exc:
        raise UnsupportedDocument(f"文件不是有效的 Office 压缩格式：{exc}") from exc
    if total > MAX_ZIP_UNCOMPRESSED_BYTES or total / compressed > MAX_ZIP_RATIO:
        raise DocumentTooLarge(
            f"文件解压后体积异常（约 {total // (1024 * 1024)}MB），已拒绝解析"
        )


def _parse_docx(path: Path) -> ParsedDocument:
    _guard_zip(path)
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - 依赖缺失时给出可执行提示
        raise UnsupportedDocument(
            "服务端缺少 python-docx，请在共享 conda 环境执行 pip install -e . 后重试"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise UnsupportedDocument(f"Word 文档无法打开（可能已加密或损坏）：{exc}") from exc

    parts: list[str] = []
    paragraphs = 0
    tables = 0
    for child in document.element.body.iterchildren():
        tag = str(child.tag).rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            paragraphs += 1
            style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                parts.append(f"{'#' * min(int(level), 6)} {text}")
            elif "list" in style:
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            tables += 1
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            rendered, _ = _rows_to_markdown(rows)
            parts.append(rendered)

    text = "\n\n".join(parts) if parts else "(文档没有可提取的文字内容)"
    return ParsedDocument(
        kind="word",
        text=text,
        meta={"paragraphs": paragraphs, "tables": tables, "format": "docx"},
    )


# ---------------------------------------------------------------------------
# Excel（.xlsx / .xlsm / .xls）
# ---------------------------------------------------------------------------


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _sheet_tsv(rows: Iterable[Iterable[Any]]) -> tuple[str, int]:
    lines: list[str] = []
    count = 0
    for row in rows:
        cells = [_format_cell(cell) for cell in row]
        if not any(cells):
            continue
        lines.append("\t".join(cells))
        count += 1
    return "\n".join(lines), count


def _parse_excel(path: Path, suffix: str, max_rows: int) -> ParsedDocument:
    if suffix == ".xls":
        return _parse_xls(path, max_rows)
    _guard_zip(path)
    return _parse_xlsx(path, max_rows)


def _parse_xlsx(path: Path, max_rows: int) -> ParsedDocument:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedDocument(
            "服务端缺少 openpyxl，请在共享 conda 环境执行 pip install -e . 后重试"
        ) from exc

    try:
        workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedDocument(
            f"Excel 工作簿无法打开（可能已加密或损坏）：{exc}"
        ) from exc

    sections: list[str] = []
    sheets: list[str] = []
    total_rows = 0
    truncated = False
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            body, row_count = _sheet_tsv(_limit_rows(rows, max_rows))
            sheets.append(worksheet.title)
            total_rows += row_count
            if max_rows > 0 and row_count >= max_rows:
                truncated = True
                body += f"\n...[工作表仅解析前 {max_rows} 行]"
            sections.append(f"## Sheet: {worksheet.title}\n\n{body or '(空表)'}")
    finally:
        with_close = getattr(workbook, "close", None)
        if callable(with_close):
            with_close()

    return ParsedDocument(
        kind="excel",
        text="\n\n".join(sections) if sections else "(工作簿没有工作表)",
        truncated=truncated,
        meta={"format": "xlsx", "sheets": sheets, "rows": total_rows},
    )


def _limit_rows(rows: Iterable[Iterable[Any]], max_rows: int) -> Iterable[Iterable[Any]]:
    if max_rows <= 0:
        yield from rows
        return
    for index, row in enumerate(rows):
        if index >= max_rows:
            return
        yield row


def _parse_xls(path: Path, max_rows: int) -> ParsedDocument:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedDocument(
            "服务端缺少 xlrd，请在共享 conda 环境执行 pip install -e . 后重试"
        ) from exc

    try:
        book = xlrd.open_workbook(str(path))
    except Exception as exc:
        raise UnsupportedDocument(
            f"旧版 Excel 无法打开（可能已加密或损坏）：{exc}"
        ) from exc

    sections: list[str] = []
    sheets: list[str] = []
    total_rows = 0
    truncated = False
    for sheet in book.sheets():
        limit = sheet.nrows if max_rows <= 0 else min(sheet.nrows, max_rows)
        rows: list[list[str]] = []
        for row_index in range(limit):
            rows.append([_xls_cell(book, sheet, row_index, col) for col in range(sheet.ncols)])
        body, row_count = _sheet_tsv(rows)
        sheets.append(sheet.name)
        total_rows += row_count
        if max_rows > 0 and sheet.nrows > max_rows:
            truncated = True
            body += f"\n...[工作表仅解析前 {max_rows} 行]"
        sections.append(f"## Sheet: {sheet.name}\n\n{body or '(空表)'}")

    return ParsedDocument(
        kind="excel",
        text="\n\n".join(sections) if sections else "(工作簿没有工作表)",
        truncated=truncated,
        meta={"format": "xls", "sheets": sheets, "rows": total_rows},
    )


def _xls_cell(book: Any, sheet: Any, row: int, col: int) -> str:
    import xlrd

    cell = sheet.cell(row, col)
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode).isoformat(sep=" ")
        except (ValueError, OverflowError):
            return str(cell.value)
    return _format_cell(cell.value)


__all__ = [
    "KIND_BY_SUFFIX",
    "LEGACY_OFFICE_HINTS",
    "MAX_TEXT_READ_BYTES",
    "SUPPORTED_SUFFIXES",
    "TRUNCATION_NOTE",
    "DocumentTooLarge",
    "ParsedDocument",
    "UnsupportedDocument",
    "decode_bytes",
    "kind_for_suffix",
    "parse_document",
    "save_extracted_text",
    "suffix_of",
]
