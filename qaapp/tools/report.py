"""``write_report_docx``：把结构化内容或 Markdown 转成 Word 报告。

只写 ``<会话目录>/outputs/reports/``；支持封面、目录域、标题层级、表格、图片、
页眉页脚页码、来源与数据质量附录。
"""

from __future__ import annotations

import re
import time
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import BaseModel, Field

from openharness.tools.base import BaseTool, ToolExecutionContext, ToolResult
from qaapp.tools.script import CONVERSATION_DIR_KEY

_BODY_FONT = "宋体"
_HEAD_FONT = "黑体"
_MAX_TABLE_ROWS = 200
_MAX_TABLE_COLS = 12


class ReportTable(BaseModel):
    title: str = ""
    columns: list[str] = Field(default_factory=list)
    rows: list[list[str]] = Field(default_factory=list)
    note: str = ""


class ReportImage(BaseModel):
    path: str
    caption: str = ""
    width_cm: float = 15.0


class ReportSection(BaseModel):
    heading: str
    level: int = 1
    paragraphs: list[str] = Field(default_factory=list)
    bullets: list[str] = Field(default_factory=list)
    tables: list[ReportTable] = Field(default_factory=list)
    images: list[ReportImage] = Field(default_factory=list)
    page_break_before: bool = False


class WriteReportDocxInput(BaseModel):
    title: str = Field(description="报告标题")
    subtitle: str = ""
    author: str = ""
    org: str = ""
    period: str = ""
    markdown: str = Field(default="", description="Markdown 全文；与 markdown_path/sections 三选一")
    markdown_path: str = Field(default="", description="会话目录内的 Markdown 文件，适合长报告")
    sections: list[ReportSection] = Field(default_factory=list, description="结构化章节；优先于 Markdown")
    sources: list[dict[str, Any]] = Field(default_factory=list, description="来源列表：label/detail/url")
    data_quality: list[str] = Field(default_factory=list, description="数据质量与局限性说明")
    file_name: str = Field(default="", description="输出文件名（不含目录），留空按标题生成")
    include_toc: bool = True
    include_page_numbers: bool = True


class WriteReportDocxTool(BaseTool):
    """生成 Word 报告。"""

    name = "write_report_docx"
    description = (
        "把分析内容生成 Word 报告（.docx），写到会话目录 outputs/reports/。\n"
        "输入可以是 sections（结构化章节/表格/图片），也可以是 markdown 或 markdown_path"
        "（适合先分段写长文，再一次性转 Word）。自动加封面、目录域、标题层级、表格样式、"
        "图片、页眉页脚页码、来源与数据质量附录。\n"
        "【要求】报告内容必须由你综合加工；数字先用 db_query 核验；不要直接把智能体原文粘进来。"
        "生成后返回路径；前端会在本轮正文下方展示结果文件卡片，用户可直接点击下载。"
    )
    input_model = WriteReportDocxInput

    def is_read_only(self, arguments: WriteReportDocxInput) -> bool:
        del arguments
        return False

    async def execute(self, arguments: WriteReportDocxInput, context: ToolExecutionContext) -> ToolResult:
        raw = (context.metadata or {}).get(CONVERSATION_DIR_KEY)
        if not raw:
            return ToolResult(output="write_report_docx 不可用：当前会话没有临时目录。", is_error=True)
        root = Path(str(raw)).resolve()
        try:
            markdown_text, source_name = self._resolve_markdown(arguments, root)
            output_path = self._output_path(root, arguments)
            warnings = self._render(
                arguments,
                markdown_text,
                source_name,
                output_path,
                root,
            )
        except ValueError as exc:
            return ToolResult(output=f"报告参数不合法：{exc}", is_error=True)
        except Exception as exc:  # noqa: BLE001 - 文档库异常统一转成可读错误
            return ToolResult(output=f"生成 Word 失败：{type(exc).__name__}: {exc}", is_error=True)
        rel = output_path.relative_to(root).as_posix()
        lines = [
            f"已生成 Word 报告：{rel}",
            f"- 大小：{output_path.stat().st_size} 字节",
            f"- 章节：{len(arguments.sections) if arguments.sections else 'Markdown 解析'}",
        ]
        if warnings:
            lines.append("- 警告：" + "；".join(warnings))
        conversation_id = str((context.metadata or {}).get("qaapp_conversation_id") or "")
        download_path = (
            f"/api/v1/conversations/{conversation_id}/files/{rel}" if conversation_id else ""
        )
        if download_path:
            lines.append(f"- 下载：{download_path}")
        else:
            lines.append("- 下载：在右侧「会话文件」中打开 outputs/reports/。")
        return ToolResult(
            output="\n".join(lines),
            metadata={
                "path": str(output_path),
                "relative_path": rel,
                "download_path": download_path,
                "warnings": warnings,
            },
        )

    # -- 输入解析 -------------------------------------------------------

    @staticmethod
    def _resolve_markdown(arguments: WriteReportDocxInput, root: Path) -> tuple[str, str]:
        if arguments.sections:
            return "", ""
        if arguments.markdown.strip():
            return arguments.markdown, "inline"
        if arguments.markdown_path.strip():
            path = WriteReportDocxTool._within(root, arguments.markdown_path)
            if not path.is_file():
                raise ValueError(f"Markdown 文件不存在：{arguments.markdown_path}")
            return path.read_text(encoding="utf-8", errors="replace"), path.name
        raise ValueError("必须提供 sections、markdown 或 markdown_path 之一")

    @staticmethod
    def _within(root: Path, relative: str) -> Path:
        candidate = (root / str(relative)).resolve()
        try:
            candidate.relative_to(root)
        except ValueError as exc:
            raise ValueError(f"路径越界：{relative}") from exc
        return candidate

    @staticmethod
    def _output_path(root: Path, arguments: WriteReportDocxInput) -> Path:
        directory = root / "outputs" / "reports"
        directory.mkdir(parents=True, exist_ok=True)
        raw_name = (arguments.file_name or arguments.title or "分析报告").strip()
        stem = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "-", raw_name)[:80] or "分析报告"
        if not stem.lower().endswith(".docx"):
            stem += ".docx"
        path = directory / stem
        if path.exists():
            path = directory / f"{path.stem}-{time.strftime('%Y%m%d-%H%M%S')}.docx"
        return path

    # -- 渲染 -----------------------------------------------------------

    def _render(
        self,
        arguments: WriteReportDocxInput,
        markdown_text: str,
        source_name: str,
        output_path: Path,
        root: Path,
    ) -> list[str]:
        warnings: list[str] = []
        doc = Document()
        self._setup_document(doc)

        self._cover(doc, arguments)
        if arguments.include_toc:
            doc.add_page_break()
            self._toc(doc)

        doc.add_page_break()
        if arguments.sections:
            self._render_sections(doc, arguments.sections, root, warnings)
        else:
            blocks = parse_markdown(markdown_text)
            self._render_blocks(doc, blocks, root, warnings)

        if arguments.data_quality:
            doc.add_heading("数据质量与局限性", level=1)
            for item in arguments.data_quality:
                self._bullet(doc, str(item))
        if arguments.sources:
            doc.add_heading("资料来源", level=1)
            for index, item in enumerate(arguments.sources, start=1):
                if isinstance(item, dict):
                    label = str(item.get("label") or item.get("title") or f"来源 {index}")
                    detail = str(item.get("detail") or item.get("url") or "")
                    self._bullet(doc, f"{label}：{detail}".rstrip("："))
                else:
                    self._bullet(doc, str(item))

        self._header_footer(doc, arguments)
        doc.save(output_path)
        return warnings

    def _setup_document(self, doc: Document) -> None:
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.6)

        normal = doc.styles["Normal"]
        normal.font.name = _BODY_FONT
        normal.font.size = Pt(12)
        _set_style_eastasia(normal, _BODY_FONT)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.first_line_indent = Cm(0.74)
        normal.paragraph_format.space_after = Pt(6)

        for level, size in ((1, 16), (2, 14), (3, 12)):
            style = doc.styles[f"Heading {level}"]
            style.font.name = _HEAD_FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            _set_style_eastasia(style, _HEAD_FONT)
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(6)

    def _cover(self, doc: Document, arguments: WriteReportDocxInput) -> None:
        for _ in range(4):
            doc.add_paragraph()
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(arguments.title)
        run.bold = True
        run.font.name = _HEAD_FONT
        run.font.size = Pt(26)
        _set_run_eastasia(run, _HEAD_FONT)
        if arguments.subtitle:
            subtitle = doc.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub_run = subtitle.add_run(arguments.subtitle)
            sub_run.font.size = Pt(14)
            sub_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            _set_run_eastasia(sub_run, _BODY_FONT)
        for _ in range(3):
            doc.add_paragraph()
        meta = [
            ("机构", arguments.org),
            ("期间", arguments.period),
            ("作者", arguments.author),
            ("生成日期", time.strftime("%Y-%m-%d")),
        ]
        for label, value in meta:
            if not value:
                continue
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(f"{label}：{value}")
            run.font.size = Pt(12)
            _set_run_eastasia(run, _BODY_FONT)

    def _toc(self, doc: Document) -> None:
        heading = doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run("目  录")
        run.bold = True
        run.font.name = _HEAD_FONT
        run.font.size = Pt(18)
        _set_run_eastasia(run, _HEAD_FONT)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = 'TOC \\o "1-3" \\h \\z \\u'
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        placeholder = OxmlElement("w:t")
        placeholder.text = "（在 Word 中右键“更新域”生成目录）"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        run._r.append(placeholder)
        run._r.append(fld_end)

    def _header_footer(self, doc: Document, arguments: WriteReportDocxInput) -> None:
        section = doc.sections[0]
        header = section.header.paragraphs[0]
        header.text = arguments.title
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            _set_run_eastasia(run, _BODY_FONT)
        if arguments.include_page_numbers:
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = footer.add_run("第 ")
            self._field(run, "PAGE")
            run2 = footer.add_run(" 页 / 共 ")
            self._field(run2, "NUMPAGES")
            footer.add_run(" 页")
            for item in footer.runs:
                item.font.size = Pt(9)

    @staticmethod
    def _field(run: Any, name: str) -> None:
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = name
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(sep)
        run._r.append(text)
        run._r.append(end)

    # -- 内容 -----------------------------------------------------------

    def _render_sections(
        self,
        doc: Document,
        sections: list[ReportSection],
        root: Path,
        warnings: list[str],
    ) -> None:
        for section in sections:
            if section.page_break_before:
                doc.add_page_break()
            if section.heading:
                doc.add_heading(section.heading, level=max(1, min(3, int(section.level))))
            for text in section.paragraphs:
                self._paragraph(doc, text)
            for text in section.bullets:
                self._bullet(doc, text)
            for table in section.tables:
                self._table(doc, table)
            for image in section.images:
                self._image(doc, image, root, warnings)

    def _render_blocks(
        self,
        doc: Document,
        blocks: list[dict[str, Any]],
        root: Path,
        warnings: list[str],
    ) -> None:
        for block in blocks:
            kind = block.get("kind")
            if kind == "heading":
                doc.add_heading(block.get("text", ""), level=int(block.get("level", 1)))
            elif kind == "paragraph":
                self._paragraph(doc, block.get("text", ""))
            elif kind == "bullet":
                self._bullet(doc, block.get("text", ""))
            elif kind == "quote":
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.74)
                run = paragraph.add_run(block.get("text", ""))
                run.italic = True
                run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
            elif kind == "code":
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.74)
                run = paragraph.add_run(block.get("text", ""))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            elif kind == "table":
                self._table(
                    doc,
                    ReportTable(
                        columns=block.get("columns", []),
                        rows=block.get("rows", []),
                    ),
                )
            elif kind == "image":
                self._image(
                    doc,
                    ReportImage(path=block.get("path", ""), caption=block.get("caption", "")),
                    root,
                    warnings,
                )
            elif kind == "hr":
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run("—" * 20)
            elif kind == "pagebreak":
                doc.add_page_break()

    def _paragraph(self, doc: Document, text: str) -> None:
        paragraph = doc.add_paragraph()
        self._add_inline(paragraph, text)

    def _bullet(self, doc: Document, text: str) -> None:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.first_line_indent = Cm(0)
        self._add_inline(paragraph, text)

    def _add_inline(self, paragraph: Any, text: str) -> None:
        """极简内联格式：**粗体**、*斜体*、`代码`、[文本](url)。"""
        pattern = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`|\[[^\]]+\]\([^)]+\))")
        position = 0
        for match in pattern.finditer(text):
            if match.start() > position:
                paragraph.add_run(text[position : match.start()])
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            elif token.startswith("["):
                label = token[1 : token.index("]")]
                paragraph.add_run(label)
            else:
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            position = match.end()
        if position < len(text):
            paragraph.add_run(text[position:])

    def _table(self, doc: Document, table: ReportTable) -> None:
        if table.title:
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = caption.add_run(table.title)
            run.bold = True
            run.font.size = Pt(10)
        columns = [str(item) for item in table.columns][:_MAX_TABLE_COLS]
        rows = table.rows[:_MAX_TABLE_ROWS]
        if not columns:
            return
        doc_table = doc.add_table(rows=1, cols=len(columns))
        doc_table.style = "Table Grid"
        doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header = doc_table.rows[0].cells
        for index, name in enumerate(columns):
            header[index].text = ""
            paragraph = header[index].paragraphs[0]
            run = paragraph.add_run(name)
            run.bold = True
            self._shade(header[index], "D9E2F3")
        for row in rows:
            cells = doc_table.add_row().cells
            for index, value in enumerate(row[: len(columns)]):
                cells[index].text = str(value if value is not None else "")
        # 表头跨页重复
        tr_pr = doc_table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
        if table.note:
            note = doc.add_paragraph()
            note.paragraph_format.first_line_indent = Cm(0)
            run = note.add_run(table.note)
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    @staticmethod
    def _shade(cell: Any, color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        tc_pr.append(shading)

    def _image(self, doc: Document, image: ReportImage, root: Path, warnings: list[str]) -> None:
        try:
            path = self._within(root, image.path)
        except ValueError:
            warnings.append(f"图片路径越界：{image.path}")
            return
        if not path.is_file():
            warnings.append(f"图片不存在：{image.path}")
            return
        try:
            doc.add_picture(str(path), width=Cm(max(5.0, min(17.0, float(image.width_cm or 15.0)))))
            if image.caption:
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption.add_run(image.caption)
                run.font.size = Pt(9)
        except Exception as exc:  # noqa: BLE001 - 单张图片失败不阻断整篇
            warnings.append(f"图片插入失败 {image.path}: {exc}")


# ---------------------------------------------------------------------------
# Markdown 子集解析
# ---------------------------------------------------------------------------


def _set_run_eastasia(run: Any, font: str) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def _set_style_eastasia(style: Any, font: str) -> None:
    """给样式设置东亚字体；不同 python-docx 版本的 rPr 可能为 None。"""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def parse_markdown(text: str) -> list[dict[str, Any]]:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[dict[str, Any]] = []
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        paragraph = " ".join(item.strip() for item in buffer if item.strip())
        if paragraph:
            blocks.append({"kind": "paragraph", "text": paragraph})
        buffer.clear()

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush()
            index += 1
            continue
        if stripped in {"<!-- pagebreak -->", "\\pagebreak", "\\newpage"}:
            flush()
            blocks.append({"kind": "pagebreak"})
            index += 1
            continue
        if stripped.startswith("```"):
            flush()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            index += 1
            blocks.append({"kind": "code", "text": "\n".join(code)})
            continue
        match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if match:
            flush()
            blocks.append({"kind": "heading", "level": len(match.group(1)), "text": match.group(2).strip()})
            index += 1
            continue
        if re.match(r"^[-*_]{3,}$", stripped):
            flush()
            blocks.append({"kind": "hr"})
            index += 1
            continue
        image = re.match(r"^!\[(.*?)\]\((.+?)\)$", stripped)
        if image:
            flush()
            blocks.append({"kind": "image", "caption": image.group(1), "path": image.group(2)})
            index += 1
            continue
        if stripped.startswith(">"):
            flush()
            blocks.append({"kind": "quote", "text": stripped.lstrip("> ").strip()})
            index += 1
            continue
        if stripped.startswith(("- ", "* ", "+ ")):
            flush()
            blocks.append({"kind": "bullet", "text": stripped[2:].strip()})
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush()
            blocks.append({"kind": "bullet", "text": re.sub(r"^\d+\.\s+", "", stripped)})
            index += 1
            continue
        if "|" in stripped and index + 1 < len(lines) and re.match(r"^\|?[\s:|-]+\|?$", lines[index + 1].strip()):
            flush()
            columns = _split_table_row(stripped)
            index += 2
            rows: list[list[str]] = []
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_split_table_row(lines[index].strip()))
                index += 1
            blocks.append({"kind": "table", "columns": columns, "rows": rows})
            continue
        buffer.append(stripped)
        index += 1
    flush()
    return blocks


def _split_table_row(line: str) -> list[str]:
    text = line.strip()
    text = text.removeprefix("|")
    text = text.removesuffix("|")
    # 支持转义竖线
    parts = re.split(r"(?<!\\)\|", text)
    return [part.replace("\\|", "|").strip() for part in parts]


__all__ = [
    "ReportImage",
    "ReportSection",
    "ReportTable",
    "WriteReportDocxInput",
    "WriteReportDocxTool",
    "parse_markdown",
]
