"""State Grid Word layout renderer."""

from __future__ import annotations

import re
import time
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from qaapp.tools.report.markdown import parse_markdown
from qaapp.tools.report.models import (
    _AUTHOR_FONT,
    _BODY_FONT,
    _BODY_INDENT_PT,
    _BODY_SIZE,
    _DEFAULT_ORG,
    _H1_FONT,
    _H3_FONT,
    _HEADING_SPECS,
    _LINE_SPACING_PT,
    _MAX_TABLE_COLS,
    _MAX_TABLE_ROWS,
    _TABLE_SIZE,
    _TITLE_FONT,
    _TITLE_SIZE,
    ReportImage,
    ReportSection,
    ReportTable,
    WriteReportDocxInput,
)
from qaapp.tools.report.styles import _set_run_eastasia, _set_style_eastasia

#: 图片段落与上下文、图注之间的间距（pt）。
_IMAGE_GAP_PT = 6.0


class ReportRendererMixin:
    """Render structured sections or Markdown into a .docx document."""

    # -- 输入解析 -------------------------------------------------------

    def _resolve_markdown(self, arguments: WriteReportDocxInput, root: Path) -> tuple[str, str]:
        if arguments.sections:
            return "", ""
        if arguments.markdown.strip():
            return arguments.markdown, "inline"
        if arguments.markdown_path.strip():
            path = self._within(root, arguments.markdown_path)
            if not path.is_file():
                raise ValueError(f"Markdown 文件不存在：{arguments.markdown_path}")
            return path.read_text(encoding="utf-8", errors="replace"), path.name
        raise ValueError("必须提供 sections、markdown 或 markdown_path 之一")

    @staticmethod
    def _within(root: Path, relative: str) -> Path:
        candidate = (root / str(relative)).resolve()
        try:
            candidate.relative_to(root)
        except ValueError as exc:
            raise ValueError(f"路径越界：{relative}") from exc
        return candidate

    @staticmethod
    def _output_path(root: Path, arguments: WriteReportDocxInput) -> Path:
        directory = root / "outputs" / "reports"
        directory.mkdir(parents=True, exist_ok=True)
        raw_name = (arguments.file_name or arguments.title or "分析报告").strip()
        stem = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "-", raw_name)[:80] or "分析报告"
        if not stem.lower().endswith(".docx"):
            stem += ".docx"
        path = directory / stem
        if path.exists():
            path = directory / f"{path.stem}-{time.strftime('%Y%m%d-%H%M%S')}.docx"
        return path

    # -- 渲染 -----------------------------------------------------------

    def _render(
        self,
        arguments: WriteReportDocxInput,
        markdown_text: str,
        source_name: str,
        output_path: Path,
        root: Path,
    ) -> list[str]:
        warnings: list[str] = []
        doc = Document()
        self._setup_document(doc)

        self._cover(doc, arguments)
        if arguments.include_toc:
            # 长报告才单独起目录页；默认标题/单位/日期与正文同页。
            doc.add_page_break()
            self._toc(doc)
            doc.add_page_break()
        else:
            separator = doc.add_paragraph()
            self._format_paragraph(separator, first_line_indent=Pt(0))

        if arguments.sections:
            self._render_sections(doc, arguments.sections, root, warnings)
        else:
            blocks = parse_markdown(markdown_text)
            self._render_blocks(doc, blocks, root, warnings)

        if arguments.data_quality:
            self._add_heading(doc, "数据质量与局限性", 1)
            for item in arguments.data_quality:
                self._bullet(doc, str(item))
        if arguments.sources:
            self._add_heading(doc, "资料来源", 1)
            for index, item in enumerate(arguments.sources, start=1):
                if isinstance(item, dict):
                    label = str(item.get("label") or item.get("title") or f"来源 {index}")
                    detail = str(item.get("detail") or item.get("url") or "")
                    self._bullet(doc, f"{label}：{detail}".rstrip("："))
                else:
                    self._bullet(doc, str(item))

        self._header_footer(doc, arguments)
        doc.save(output_path)
        return warnings

    @staticmethod
    def _today_cn() -> str:
        now = datetime.now(ZoneInfo("Asia/Shanghai"))
        return f"{now.year}年{now.month}月{now.day}日"

    @staticmethod
    def _format_paragraph(
        paragraph: Any,
        *,
        alignment: Any = None,
        first_line_indent: Any = None,
    ) -> None:
        pf = paragraph.paragraph_format
        if alignment is not None:
            pf.alignment = alignment
        if first_line_indent is not None:
            pf.first_line_indent = first_line_indent
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(_LINE_SPACING_PT)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)

    @staticmethod
    def _format_run(run: Any, font: str, size: float, bold: bool | None = None) -> None:
        run.font.name = font
        run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        _set_run_eastasia(run, font)

    def _add_heading(self, doc: Document, text: str, level: int) -> Any:
        """Add a 1–4 level heading with the State Grid font rules."""
        level = max(1, min(4, int(level)))
        paragraph = doc.add_heading(str(text), level=level)
        self._format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_indent=Pt(0),
        )
        paragraph.paragraph_format.keep_with_next = True
        font, bold = _HEADING_SPECS[level]
        for run in paragraph.runs:
            self._format_run(run, font, _BODY_SIZE, bold)
        return paragraph

    def _setup_document(self, doc: Document) -> None:
        section = doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        # 国网公文格式：上下 35mm，左右 28mm。
        section.top_margin = Mm(35)
        section.bottom_margin = Mm(35)
        section.left_margin = Mm(28)
        section.right_margin = Mm(28)
        normal = doc.styles["Normal"]
        normal.font.name = _BODY_FONT
        normal.font.size = Pt(_BODY_SIZE)
        normal.font.bold = False
        _set_style_eastasia(normal, _BODY_FONT)
        normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        normal.paragraph_format.line_spacing = Pt(_LINE_SPACING_PT)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.first_line_indent = Pt(0)

        for level, (font, bold) in _HEADING_SPECS.items():
            style = doc.styles[f"Heading {level}"]
            style.font.name = font
            style.font.size = Pt(_BODY_SIZE)
            style.font.bold = bold
            style.font.color.rgb = RGBColor(0, 0, 0)
            _set_style_eastasia(style, font)
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            style.paragraph_format.line_spacing = Pt(_LINE_SPACING_PT)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.first_line_indent = Pt(0)
            style.paragraph_format.keep_with_next = True

    def _cover(self, doc: Document, arguments: WriteReportDocxInput) -> None:
        title = doc.add_paragraph()
        self._format_paragraph(
            title,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            first_line_indent=Pt(0),
        )
        run = title.add_run(arguments.title)
        self._format_run(run, _TITLE_FONT, _TITLE_SIZE, True)

        subtitle = arguments.subtitle.strip() or arguments.period.strip()
        if subtitle:
            paragraph = doc.add_paragraph()
            self._format_paragraph(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                first_line_indent=Pt(0),
            )
            run = paragraph.add_run(subtitle)
            self._format_run(run, _AUTHOR_FONT, _BODY_SIZE, False)

        org = arguments.org.strip()
        author = arguments.author.strip()
        lines = [value for value in (org, author) if value and value != org]
        if org:
            lines.insert(0, org)
        if not lines:
            lines.append(_DEFAULT_ORG)
        lines.append(arguments.date.strip() or self._today_cn())
        for value in lines:
            paragraph = doc.add_paragraph()
            self._format_paragraph(
                paragraph,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                first_line_indent=Pt(0),
            )
            run = paragraph.add_run(value)
            self._format_run(run, _AUTHOR_FONT, _BODY_SIZE, True)

    def _toc(self, doc: Document) -> None:
        heading = doc.add_paragraph()
        self._format_paragraph(
            heading,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
            first_line_indent=Pt(0),
        )
        run = heading.add_run("目  录")
        self._format_run(run, _H1_FONT, _BODY_SIZE, False)

        paragraph = doc.add_paragraph()
        self._format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_indent=Pt(0),
        )
        run = paragraph.add_run()
        self._format_run(run, _BODY_FONT, _BODY_SIZE, None)
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = 'TOC \\o "1-4" \\h \\z \\u'
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        placeholder = OxmlElement("w:t")
        placeholder.text = "（在 Word 中右键“更新域”生成目录）"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_sep)
        run._r.append(placeholder)
        run._r.append(fld_end)

    def _header_footer(self, doc: Document, arguments: WriteReportDocxInput) -> None:
        if not arguments.include_page_numbers:
            return
        section = doc.sections[0]
        # 只有单独封面页时才隐藏首页页码；默认标题与正文同页，首页正常编号。
        section.different_first_page_header_footer = bool(arguments.include_toc)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("第 ")
        self._field(run, "PAGE")
        run2 = footer.add_run(" 页 / 共 ")
        self._field(run2, "NUMPAGES")
        footer.add_run(" 页")
        for item in footer.runs:
            item.font.name = "宋体"
            item.font.size = Pt(9)
            _set_run_eastasia(item, "宋体")

    @staticmethod
    def _field(run: Any, name: str) -> None:
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = name
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.append(begin)
        run._r.append(instr)
        run._r.append(sep)
        run._r.append(text)
        run._r.append(end)

    # -- 内容 -----------------------------------------------------------

    def _render_sections(
        self,
        doc: Document,
        sections: list[ReportSection],
        root: Path,
        warnings: list[str],
    ) -> None:
        for section in sections:
            if section.page_break_before:
                doc.add_page_break()
            if section.heading:
                self._add_heading(doc, section.heading, section.level)
            for text in section.paragraphs:
                self._paragraph(doc, text)
            for text in section.bullets:
                self._bullet(doc, text)
            for table in section.tables:
                self._table(doc, table)
            for image in section.images:
                self._image(doc, image, root, warnings)

    def _render_blocks(
        self,
        doc: Document,
        blocks: list[dict[str, Any]],
        root: Path,
        warnings: list[str],
    ) -> None:
        for block in blocks:
            kind = block.get("kind")
            if kind == "heading":
                self._add_heading(doc, block.get("text", ""), int(block.get("level", 1)))
            elif kind == "paragraph":
                self._paragraph(doc, block.get("text", ""))
            elif kind == "bullet":
                self._bullet(doc, block.get("text", ""))
            elif kind == "quote":
                paragraph = doc.add_paragraph()
                self._format_paragraph(
                    paragraph,
                    alignment=WD_ALIGN_PARAGRAPH.LEFT,
                    first_line_indent=Pt(0),
                )
                paragraph.paragraph_format.left_indent = Cm(0.74)
                run = paragraph.add_run(block.get("text", ""))
                self._format_run(run, _BODY_FONT, _BODY_SIZE, False)
            elif kind == "code":
                paragraph = doc.add_paragraph()
                self._format_paragraph(
                    paragraph,
                    alignment=WD_ALIGN_PARAGRAPH.LEFT,
                    first_line_indent=Pt(0),
                )
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                run = paragraph.add_run(block.get("text", ""))
                self._format_run(run, _BODY_FONT, _TABLE_SIZE, False)
            elif kind == "table":
                self._table(
                    doc,
                    ReportTable(
                        columns=block.get("columns", []),
                        rows=block.get("rows", []),
                    ),
                )
            elif kind == "image":
                self._image(
                    doc,
                    ReportImage(path=block.get("path", ""), caption=block.get("caption", "")),
                    root,
                    warnings,
                )
            elif kind == "hr":
                paragraph = doc.add_paragraph()
                self._format_paragraph(
                    paragraph,
                    alignment=WD_ALIGN_PARAGRAPH.CENTER,
                    first_line_indent=Pt(0),
                )
                run = paragraph.add_run("—" * 20)
                self._format_run(run, _BODY_FONT, _BODY_SIZE, False)
            elif kind == "pagebreak":
                doc.add_page_break()

    def _paragraph(self, doc: Document, text: str) -> None:
        paragraph = doc.add_paragraph()
        self._format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
            first_line_indent=Pt(_BODY_INDENT_PT),
        )
        self._add_inline(paragraph, text)

    def _bullet(self, doc: Document, text: str) -> None:
        paragraph = doc.add_paragraph(style="List Bullet")
        self._format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_indent=Pt(0),
        )
        paragraph.paragraph_format.left_indent = Cm(0.74)
        self._add_inline(paragraph, text)

    def _add_inline(self, paragraph: Any, text: str) -> None:
        """极简内联格式：**粗体**、*斜体*、`代码`、[文本](url)。"""
        pattern = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`|\[[^\]]+\]\([^)]+\))")
        position = 0
        for match in pattern.finditer(text):
            if match.start() > position:
                paragraph.add_run(text[position : match.start()])
            token = match.group(0)
            if token.startswith("**"):
                run = paragraph.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(10)
            elif token.startswith("["):
                label = token[1 : token.index("]")]
                paragraph.add_run(label)
            else:
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            position = match.end()
        if position < len(text):
            paragraph.add_run(text[position:])
        for run in paragraph.runs:
            self._format_run(run, _BODY_FONT, _BODY_SIZE, None)

    def _table(self, doc: Document, table: ReportTable) -> None:
        if table.title:
            caption = doc.add_paragraph()
            self._format_paragraph(
                caption,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                first_line_indent=Pt(0),
            )
            run = caption.add_run(table.title)
            self._format_run(run, _H3_FONT, _TABLE_SIZE, True)
        columns = [str(item) for item in table.columns][:_MAX_TABLE_COLS]
        rows = table.rows[:_MAX_TABLE_ROWS]
        if not columns:
            return
        doc_table = doc.add_table(rows=1, cols=len(columns))
        doc_table.style = "Table Grid"
        doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header = doc_table.rows[0].cells
        for index, name in enumerate(columns):
            self._table_cell(header[index], name, bold=True, font=_H3_FONT)
            self._shade(header[index], "F2F2F2")
        for row in rows:
            cells = doc_table.add_row().cells
            for index, value in enumerate(row[: len(columns)]):
                self._table_cell(
                    cells[index],
                    str(value if value is not None else ""),
                    bold=False,
                    font=_BODY_FONT,
                )
        # 表头跨页重复
        tr_pr = doc_table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
        if table.note:
            note = doc.add_paragraph()
            self._format_paragraph(
                note,
                alignment=WD_ALIGN_PARAGRAPH.LEFT,
                first_line_indent=Pt(0),
            )
            run = note.add_run(table.note)
            self._format_run(run, _BODY_FONT, 9, False)

    def _table_cell(self, cell: Any, text: str, *, bold: bool, font: str) -> None:
        cell.text = ""
        paragraph = cell.paragraphs[0]
        self._format_paragraph(
            paragraph,
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_indent=Pt(0),
        )
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        run = paragraph.add_run(str(text))
        self._format_run(run, font, _TABLE_SIZE, bold)

    @staticmethod
    def _shade(cell: Any, color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        tc_pr.append(shading)

    def _image(self, doc: Document, image: ReportImage, root: Path, warnings: list[str]) -> None:
        try:
            path = self._within(root, image.path)
        except ValueError:
            warnings.append(f"图片路径越界：{image.path}")
            return
        if not path.is_file():
            warnings.append(f"图片不存在：{image.path}")
            return
        try:
            # A4 正文宽度 = 21 - 2.8 - 2.8 = 15.4cm，图片不得越界。
            inline_shape = doc.add_picture(
                str(path),
                width=Cm(max(5.0, min(15.4, float(image.width_cm or 15.0)))),
            )
            picture_paragraph = doc.paragraphs[-1]
            self._format_paragraph(
                picture_paragraph,
                alignment=WD_ALIGN_PARAGRAPH.CENTER,
                first_line_indent=Pt(0),
            )
            # 图片段落不能用正文的 29pt 固定行距，否则图片会被裁成一条并压住正文；
            # 也不能按图高预留行高，否则 Word/WPS 会在图片下方再空出一整张图的高度。
            # 这里用单倍行距 + 1pt 段落标记：浮动图片只由“上下环绕”占位，空段落几乎不占高。
            picture_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            picture_paragraph.paragraph_format.space_before = Pt(_IMAGE_GAP_PT)
            picture_paragraph.paragraph_format.space_after = Pt(0)
            self._compact_paragraph_mark(picture_paragraph)
            # 改为“上下环绕”浮动图片，文字只出现在图片上方和下方。
            self._wrap_top_and_bottom(inline_shape)
            if image.caption:
                # 图与图注不拆页：图注若放不下，整块留到下一页，避免“图在这页、图注在下页”。
                picture_paragraph.paragraph_format.keep_with_next = True
                caption = doc.add_paragraph()
                self._format_paragraph(
                    caption,
                    alignment=WD_ALIGN_PARAGRAPH.CENTER,
                    first_line_indent=Pt(0),
                )
                # 图注用单倍行距，避免 29pt 行距把图注推离图片。
                caption.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                caption.paragraph_format.space_before = Pt(_IMAGE_GAP_PT)
                caption.paragraph_format.space_after = Pt(0)
                run = caption.add_run(image.caption)
                self._format_run(run, _BODY_FONT, _TABLE_SIZE, False)
        except Exception as exc:  # noqa: BLE001 - 单张图片失败不阻断整篇
            warnings.append(f"图片插入失败 {image.path}: {exc}")

    @staticmethod
    def _compact_paragraph_mark(paragraph: Any) -> None:
        """把空段落的段落标记压到 1pt，并关闭网格对齐。

        图片段落的正文只有一张浮动图片；段落标记若仍是三号字（且吸附文档网格），
        会在图片下方撑出一整行空白。压小标记字号后，这段空行基本不占高度。
        """
        p_pr = paragraph._p.get_or_add_pPr()

        snap_to_grid = p_pr.find(qn("w:snapToGrid"))
        if snap_to_grid is None:
            snap_to_grid = OxmlElement("w:snapToGrid")
            snap_to_grid.set(qn("w:val"), "0")
            p_pr.insert_element_before(
                snap_to_grid,
                "w:spacing",
                "w:ind",
                "w:contextualSpacing",
                "w:mirrorIndents",
                "w:suppressOverlap",
                "w:jc",
                "w:textDirection",
                "w:textAlignment",
                "w:textboxTightWrap",
                "w:outlineLvl",
                "w:divId",
                "w:cnfStyle",
                "w:rPr",
                "w:sectPr",
                "w:pPrChange",
            )

        r_pr = p_pr.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = OxmlElement("w:rPr")
            p_pr.insert_element_before(r_pr, "w:sectPr", "w:pPrChange")
        for tag in ("w:sz", "w:szCs"):
            if r_pr.find(qn(tag)) is None:
                size = OxmlElement(tag)
                size.set(qn("w:val"), "2")  # 半磅制：2 = 1pt
                r_pr.append(size)

    @staticmethod
    def _wrap_top_and_bottom(shape: Any, *, gap_emu: int = 76200) -> None:
        """把内嵌图片改写成锚定图片，环绕方式为“上下环绕”。

        python-docx 只支持内嵌（inline）图片；内嵌图片位于行内，一旦所在段落
        使用固定行距，Word 会按行高裁切图片，导致图片与正文重叠。这里改写
        ``wp:inline`` 为 ``wp:anchor`` + ``wp:wrapTopAndBottom``：图片按列居中、
        上下文字自动让位，既有间距也不会串行。
        """
        inline = shape._inline
        parent = inline.getparent()
        if parent is None:  # pragma: no cover - 正常文档不会出现
            return

        anchor = OxmlElement("wp:anchor")
        for name, value in (
            ("distT", str(gap_emu)),
            ("distB", str(gap_emu)),
            ("distL", "0"),
            ("distR", "0"),
            ("simplePos", "0"),
            ("relativeHeight", "251658240"),
            ("behindDoc", "0"),
            ("locked", "0"),
            ("layoutInCell", "1"),
            ("allowOverlap", "0"),
        ):
            anchor.set(name, value)

        simple_pos = OxmlElement("wp:simplePos")
        simple_pos.set("x", "0")
        simple_pos.set("y", "0")
        anchor.append(simple_pos)

        position_h = OxmlElement("wp:positionH")
        position_h.set("relativeFrom", "column")
        align = OxmlElement("wp:align")
        align.text = "center"
        position_h.append(align)
        anchor.append(position_h)

        position_v = OxmlElement("wp:positionV")
        position_v.set("relativeFrom", "paragraph")
        # 用 align=top 而不是 posOffset=0：部分阅读器（LibreOffice/WPS）会把
        # 段落相对偏移 0 解释成“贴到段落底部”，从而把图片排到前一段之前。
        vertical = OxmlElement("wp:align")
        vertical.text = "top"
        position_v.append(vertical)
        anchor.append(position_v)

        extent = inline.find(qn("wp:extent"))
        if extent is not None:
            anchor.append(deepcopy(extent))

        effect_extent = inline.find(qn("wp:effectExtent"))
        if effect_extent is not None:
            anchor.append(deepcopy(effect_extent))
        else:
            effect_extent = OxmlElement("wp:effectExtent")
            for name in ("l", "t", "r", "b"):
                effect_extent.set(name, "0")
            anchor.append(effect_extent)

        anchor.append(OxmlElement("wp:wrapTopAndBottom"))

        doc_pr = inline.find(qn("wp:docPr"))
        if doc_pr is not None:
            anchor.append(deepcopy(doc_pr))

        frame_pr = inline.find(qn("wp:cNvGraphicFramePr"))
        if frame_pr is not None:
            anchor.append(deepcopy(frame_pr))
        else:
            frame_pr = OxmlElement("wp:cNvGraphicFramePr")
            locks = OxmlElement("a:graphicFrameLocks")
            locks.set("noChangeAspect", "1")
            frame_pr.append(locks)
            anchor.append(frame_pr)

        graphic = inline.find(qn("a:graphic"))
        if graphic is None:  # pragma: no cover - 缺少图形节点时保持内嵌
            return
        anchor.append(deepcopy(graphic))

        parent.replace(inline, anchor)
